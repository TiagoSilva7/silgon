import argparse
import concurrent.futures
import json
import os
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from typing import Any, Dict, Iterable, List, Optional, Tuple

import requests
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry


DEFAULT_TIMEOUT = 30
KEYWORDS = ("CPF", "CNPJ")


# ==========================================================
# PREENCHA AQUI (opcional)
# REST 2024 OAS 3.0 - informe URL base da Library (com ou sem /api)
# Exemplo: http://servidor:8080/MicroStrategyLibrary
# ==========================================================
MSTR_BASE_URL = "http://10.14.203.158:8080/SEFADEVLIB/api"
MSTR_USERNAME = "31071655817"
MSTR_PASSWORD = "Tera!7777"
MSTR_PROJECT_ID = "239BC40211E8B1E3626F0080EF95EED2"
MSTR_PROJECT_NAME = ""


@dataclass
class ApiContext:
    base_url: str
    timeout: int
    verify_ssl: bool


class MstrClient:
    def __init__(self, ctx: ApiContext):
        self.ctx = ctx
        self.session = self._build_session()

    def _build_session(self) -> requests.Session:
        session = requests.Session()
        session.verify = self.ctx.verify_ssl
        session.headers.update({"Accept": "application/json"})

        retry = Retry(
            total=3,
            read=3,
            connect=3,
            backoff_factor=0.5,
            status_forcelist=[429, 500, 502, 503, 504],
            allowed_methods=frozenset(["GET", "POST"]),
        )
        adapter = HTTPAdapter(pool_connections=40, pool_maxsize=40, max_retries=retry)
        session.mount("http://", adapter)
        session.mount("https://", adapter)
        return session

    def clone(self) -> "MstrClient":
        cloned = MstrClient(self.ctx)
        cloned.session.headers.update(self.session.headers)
        cloned.session.cookies.update(self.session.cookies)
        return cloned

    def _url(self, path: str) -> str:
        return f"{self.ctx.base_url.rstrip('/')}/{path.lstrip('/')}"

    def get(self, path: str, **kwargs: Any) -> requests.Response:
        return self.session.get(self._url(path), timeout=self.ctx.timeout, **kwargs)

    def post(self, path: str, **kwargs: Any) -> requests.Response:
        return self.session.post(self._url(path), timeout=self.ctx.timeout, **kwargs)

    def login(self, username: str, password: str) -> None:
        payload = {"username": username, "password": password, "loginMode": 1}
        response = self.post("/api/auth/login", json=payload)
        response.raise_for_status()

        auth_token = response.headers.get("X-MSTR-AuthToken")
        if not auth_token:
            raise RuntimeError("Não foi possível obter X-MSTR-AuthToken no login.")

        self.session.headers.update(
            {"X-MSTR-AuthToken": auth_token, "Content-Type": "application/json"}
        )

    def logout(self) -> None:
        try:
            self.post("/api/auth/logout")
        except requests.RequestException:
            pass

    def set_project(self, project_id: str) -> None:
        self.session.headers.update({"X-MSTR-ProjectID": project_id})


def normalize_base_url(raw_base_url: str) -> str:
    normalized = raw_base_url.strip().rstrip("/")
    if normalized.lower().endswith("/api"):
        normalized = normalized[:-4]
    return normalized


def first_not_empty(*values: Optional[str]) -> Optional[str]:
    for value in values:
        if value:
            return value
    return None


def contains_keyword(value: Optional[str], keywords: Iterable[str] = KEYWORDS) -> bool:
    text = (value or "").upper()
    return any(keyword in text for keyword in keywords)


def normalize_name(value: Optional[str]) -> str:
    return re.sub(r"[^A-Z0-9]", "", (value or "").upper())


def find_project_id(client: MstrClient, project_id: Optional[str], project_name: Optional[str]) -> str:
    if project_id:
        return project_id

    response = client.get("/api/projects")
    response.raise_for_status()
    payload = response.json()
    projects = payload.get("projects", []) if isinstance(payload, dict) else []

    if not projects:
        raise RuntimeError("Nenhum projeto foi retornado por /api/projects.")

    if project_name:
        target = project_name.strip().lower()
        for project in projects:
            if str(project.get("name", "")).strip().lower() == target:
                return str(project.get("id"))
        raise RuntimeError(f"Projeto '{project_name}' não encontrado.")

    first_project_id = projects[0].get("id")
    if not first_project_id:
        raise RuntimeError("Não foi possível determinar o ID do projeto padrão.")
    return str(first_project_id)


def get_search_results_paginated(
    client: MstrClient,
    object_type: int,
    limit: int = 10000,
) -> List[Dict[str, Any]]:
    offset = 0
    total_items: Optional[int] = None
    all_items: List[Dict[str, Any]] = []

    while True:
        response = client.get(f"/api/searches/results?type={object_type}&limit={limit}&offset={offset}")
        if response.status_code >= 400:
            raise RuntimeError(
                f"Falha em /api/searches/results type={object_type} offset={offset}: HTTP {response.status_code}"
            )

        payload = response.json()
        if not isinstance(payload, dict):
            break

        if total_items is None:
            total_items = int(payload.get("totalItems", 0))

        page_items = payload.get("result", [])
        if not isinstance(page_items, list) or not page_items:
            break

        all_items.extend(item for item in page_items if isinstance(item, dict))
        offset += limit

        if total_items is not None and offset >= total_items:
            break

    return all_items


def get_attribute_datatype_from_object(client: MstrClient, attribute_id: str) -> Optional[str]:
    try:
        response = client.get(f"/api/objects/{attribute_id}?type=12")
        if response.status_code >= 400:
            return None
        payload = response.json()
        if not isinstance(payload, dict):
            return None

        for key in ("dataType", "datatype", "data_type", "typeName"):
            value = payload.get(key)
            if isinstance(value, str) and value.strip():
                return value.strip()

        return None
    except Exception:
        return None


def get_attributes_2024(
    client: MstrClient,
    workers: int,
    resolve_datatype: bool,
) -> List[Dict[str, Optional[str]]]:
    items = get_search_results_paginated(client, object_type=12, limit=10000)

    attributes: List[Dict[str, Optional[str]]] = []
    missing_dtype_indexes: List[Tuple[int, str]] = []

    for item in items:
        attribute_id = str(item.get("id", "")).strip() or None
        attribute_name = str(item.get("name", "")).strip() or None

        row = {
            "id": attribute_id,
            "name": attribute_name,
            "data_type": None,
            "subtype": item.get("subtype"),
            "extType": item.get("extType"),
        }
        attributes.append(row)

        if resolve_datatype and attribute_id:
            missing_dtype_indexes.append((len(attributes) - 1, attribute_id))

    if resolve_datatype and missing_dtype_indexes:
        print(f"Resolvendo data_type de {len(missing_dtype_indexes)} atributos (pode demorar)...")

        def resolve_one(attribute_id: str) -> Optional[str]:
            return get_attribute_datatype_from_object(client.clone(), attribute_id)

        completed = 0
        with concurrent.futures.ThreadPoolExecutor(max_workers=max(1, workers)) as executor:
            futures = {
                executor.submit(resolve_one, attribute_id): (index, attribute_id)
                for index, attribute_id in missing_dtype_indexes
            }
            for future in concurrent.futures.as_completed(futures):
                completed += 1
                index, _ = futures[future]
                try:
                    dtype = future.result()
                    if dtype:
                        attributes[index]["data_type"] = dtype
                except Exception:
                    pass

                if completed % 500 == 0 or completed == len(missing_dtype_indexes):
                    print(f"Atributos detalhados: {completed}/{len(missing_dtype_indexes)}")

    attributes.sort(key=lambda item: (item.get("name") or "").lower())
    return attributes


def get_reports_2024(client: MstrClient) -> List[Dict[str, Any]]:
    items = get_search_results_paginated(client, object_type=3, limit=10000)

    reports: List[Dict[str, Any]] = []
    for item in items:
        report_id = str(item.get("id", "")).strip() or None
        report_name = str(item.get("name", "")).strip() or None
        if not report_id or not report_name:
            continue

        reports.append(
            {
                "id": report_id,
                "name": report_name,
                "type": item.get("type"),
                "subtype": item.get("subtype"),
                "extType": item.get("extType"),
            }
        )

    return reports


def get_report_definition_2024(client: MstrClient, report_id: str) -> Optional[Dict[str, Any]]:
    try:
        response = client.get(f"/api/reports/{report_id}")
        if response.status_code >= 400:
            return None
        payload = response.json()
        if not isinstance(payload, dict):
            return None

        result = payload.get("result")
        if isinstance(result, dict):
            definition = result.get("definition")
            if isinstance(definition, dict):
                return definition

        return None
    except Exception:
        return None


def extract_columns_from_definition(definition: Dict[str, Any]) -> List[Dict[str, Optional[str]]]:
    template = definition.get("template") if isinstance(definition, dict) else None
    if not isinstance(template, dict):
        return []

    collected: List[Tuple[str, Optional[str], Optional[str]]] = []

    for axis in ("rows", "columns"):
        axis_items = template.get(axis, [])
        if not isinstance(axis_items, list):
            continue

        for item in axis_items:
            if not isinstance(item, dict):
                continue

            name = str(item.get("name", "")).strip() or None
            if not name:
                continue

            dtype = None
            column_alias = None
            forms = item.get("forms", [])
            if isinstance(forms, list):
                for form in forms:
                    if not isinstance(form, dict):
                        continue
                    column_alias = str(form.get("name", "")).strip() or column_alias
                    dtype = first_not_empty(
                        str(form.get("dataType", "")).strip() or None,
                        str(form.get("baseFormType", "")).strip() or None,
                    )
                    if dtype:
                        break

            collected.append((name, dtype, column_alias))

    metrics = template.get("metrics", [])
    if isinstance(metrics, list):
        for metric in metrics:
            if not isinstance(metric, dict):
                continue
            name = str(metric.get("name", "")).strip() or None
            if name:
                collected.append((name, None, None))

    dedup: Dict[Tuple[str, str, str], Dict[str, Optional[str]]] = {}
    for name, dtype, column_alias in collected:
        key = (name.lower(), (dtype or "").lower(), (column_alias or "").lower())
        dedup[key] = {"name": name, "data_type": dtype, "column_alias": column_alias}

    return list(dedup.values())


def filter_keyword_columns(columns: List[Dict[str, Optional[str]]]) -> List[Dict[str, Optional[str]]]:
    return [col for col in columns if contains_keyword(col.get("name"))]


def enrich_keyword_attributes_from_report_columns(
    attributes_keyword: List[Dict[str, Optional[str]]],
    freeform_reports: List[Dict[str, Any]],
) -> None:
    by_norm_name: Dict[str, Dict[str, Optional[str]]] = {}

    for report in freeform_reports:
        for column in report.get("columns", []):
            col_name = column.get("name")
            norm = normalize_name(col_name)
            if not norm:
                continue

            existing = by_norm_name.get(norm)
            if existing is None:
                by_norm_name[norm] = {
                    "column_alias": column.get("column_alias"),
                    "column_data_type": column.get("data_type"),
                    "source_report_id": report.get("id"),
                    "source_report_name": report.get("name"),
                }
                continue

            # Prioriza registros com data type preenchido.
            if not existing.get("column_data_type") and column.get("data_type"):
                by_norm_name[norm] = {
                    "column_alias": column.get("column_alias"),
                    "column_data_type": column.get("data_type"),
                    "source_report_id": report.get("id"),
                    "source_report_name": report.get("name"),
                }

    for attribute in attributes_keyword:
        norm_attr_name = normalize_name(attribute.get("name"))
        matched = by_norm_name.get(norm_attr_name)
        if not matched:
            continue

        attribute["column_alias"] = matched.get("column_alias")
        attribute["column_alias_data_type"] = matched.get("column_data_type")
        attribute["column_alias_source_report_id"] = matched.get("source_report_id")
        attribute["column_alias_source_report_name"] = matched.get("source_report_name")

        if not attribute.get("data_type") and matched.get("column_data_type"):
            attribute["data_type"] = matched.get("column_data_type")


def is_freeform_report_metadata(report: Dict[str, Any]) -> bool:
    # REST 2024: metadado de listagem traz extType.
    # No ambiente validado, relatórios Freeform SQL aparecem com extType=3.
    return report.get("extType") == 3


def analyze_one_report_2024(client: MstrClient, report: Dict[str, Any]) -> Optional[Dict[str, Any]]:
    if not is_freeform_report_metadata(report):
        return None

    definition = get_report_definition_2024(client, report["id"])
    if not definition:
        return None

    columns = extract_columns_from_definition(definition)
    matched_columns = filter_keyword_columns(columns)
    if not matched_columns:
        return None

    return {
        "id": report["id"],
        "name": report["name"],
        "extType": report.get("extType"),
        "subtype": report.get("subtype"),
        "columns": matched_columns,
    }


def get_freeform_reports_with_cpf_cnpj_2024(
    client: MstrClient,
    workers: int,
    include_report_ids: List[str],
) -> List[Dict[str, Any]]:
    reports = get_reports_2024(client)

    reports_by_id: Dict[str, Dict[str, Any]] = {report["id"]: report for report in reports}

    # IDs informados manualmente entram como candidatos freeform, mesmo que metadata venha diferente.
    for report_id in include_report_ids:
        if not report_id:
            continue
        if report_id in reports_by_id:
            reports_by_id[report_id]["extType"] = 3
        else:
            reports_by_id[report_id] = {
                "id": report_id,
                "name": f"(informado manualmente) {report_id}",
                "type": 3,
                "subtype": 768,
                "extType": 3,
            }

    freeform_candidates = [report for report in reports_by_id.values() if report.get("extType") == 3]

    total = len(freeform_candidates)
    print(f"Total de relatórios (paginado): {len(reports_by_id)}")
    print(f"Candidatos Freeform (extType=3): {total}")
    if total == 0:
        return []

    results: List[Dict[str, Any]] = []

    def worker(report: Dict[str, Any]) -> Optional[Dict[str, Any]]:
        return analyze_one_report_2024(client.clone(), report)

    completed = 0
    with concurrent.futures.ThreadPoolExecutor(max_workers=max(1, workers)) as executor:
        futures = [executor.submit(worker, report) for report in freeform_candidates]
        for future in concurrent.futures.as_completed(futures):
            completed += 1
            try:
                item = future.result()
                if item:
                    results.append(item)
            except Exception:
                pass

            if completed % 100 == 0 or completed == total:
                print(f"Relatórios processados: {completed}/{total} | Encontrados: {len(results)}")

    results.sort(key=lambda item: item["name"].lower())
    return results


def save_excel(
    file_path: str,
    attributes: List[Dict[str, Optional[str]]],
    attributes_keyword: List[Dict[str, Optional[str]]],
    freeform_reports: List[Dict[str, Any]],
) -> None:
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise RuntimeError("Pacote 'openpyxl' não encontrado. Instale com: pip install openpyxl") from exc

    workbook = Workbook()

    ws_attr = workbook.active
    ws_attr.title = "Atributos_Todos"
    ws_attr.append(["attribute_id", "attribute_name", "data_type", "subtype", "extType"])
    for attribute in attributes:
        ws_attr.append(
            [
                attribute.get("id"),
                attribute.get("name"),
                attribute.get("data_type"),
                attribute.get("subtype"),
                attribute.get("extType"),
            ]
        )

    ws_attr_kw = workbook.create_sheet("Atributos_CPF_CNPJ")
    ws_attr_kw.append(
        [
            "attribute_id",
            "attribute_name",
            "data_type",
            "column_alias",
            "column_alias_data_type",
            "column_alias_source_report_id",
            "column_alias_source_report_name",
            "subtype",
            "extType",
        ]
    )
    for attribute in attributes_keyword:
        ws_attr_kw.append(
            [
                attribute.get("id"),
                attribute.get("name"),
                attribute.get("data_type"),
                attribute.get("column_alias"),
                attribute.get("column_alias_data_type"),
                attribute.get("column_alias_source_report_id"),
                attribute.get("column_alias_source_report_name"),
                attribute.get("subtype"),
                attribute.get("extType"),
            ]
        )

    ws_reports = workbook.create_sheet("Relatorios_Freeform")
    ws_reports.append(
        [
            "report_id",
            "report_name",
            "extType",
            "subtype",
            "column_name",
            "column_alias",
            "column_data_type",
        ]
    )
    for report in freeform_reports:
        for column in report.get("columns", []):
            ws_reports.append(
                [
                    report.get("id"),
                    report.get("name"),
                    report.get("extType"),
                    report.get("subtype"),
                    column.get("name"),
                    column.get("column_alias"),
                    column.get("data_type"),
                ]
            )

    workbook.save(file_path)


def parse_include_report_ids(raw_value: str) -> List[str]:
    if not raw_value:
        return []
    return [part.strip() for part in raw_value.split(",") if part.strip()]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "MicroStrategy REST 2024 (OAS 3.0): lista atributos e relatórios Freeform SQL "
            "com CPF/CNPJ em nomes de colunas."
        )
    )
    parser.add_argument(
        "--base-url",
        default=(os.getenv("MSTR_BASE_URL") or MSTR_BASE_URL or None),
        help="URL base da Library (com ou sem /api), ex: http://host:8080/MicroStrategyLibrary",
    )
    parser.add_argument(
        "--username",
        default=(os.getenv("MSTR_USERNAME") or MSTR_USERNAME or None),
        help="Usuário MicroStrategy",
    )
    parser.add_argument(
        "--password",
        default=(os.getenv("MSTR_PASSWORD") or MSTR_PASSWORD or None),
        help="Senha MicroStrategy",
    )
    parser.add_argument(
        "--project-id",
        default=(os.getenv("MSTR_PROJECT_ID") or MSTR_PROJECT_ID or None),
        help="ID do projeto",
    )
    parser.add_argument(
        "--project-name",
        default=(os.getenv("MSTR_PROJECT_NAME") or MSTR_PROJECT_NAME or None),
        help="Nome do projeto (se não informar ID)",
    )
    parser.add_argument("--timeout", type=int, default=DEFAULT_TIMEOUT, help="Timeout de requisição em segundos")
    parser.add_argument("--workers", type=int, default=12, help="Número de workers paralelos")
    parser.add_argument(
        "--resolve-attribute-datatype",
        action="store_true",
        help="Tenta obter data_type de atributos via /api/objects/{id}?type=12 (mais lento)",
    )
    parser.add_argument(
        "--include-report-ids",
        default="",
        help="IDs de relatórios (separados por vírgula) para forçar inclusão na análise",
    )
    parser.add_argument("--insecure", action="store_true", help="Desabilita validação SSL")
    parser.add_argument("--output", default="resultado_mstr_cpf_cnpj.json", help="Arquivo de saída JSON")
    parser.add_argument("--excel-output", default="resultado_mstr_cpf_cnpj.xlsx", help="Arquivo de saída Excel")
    return parser.parse_args()


def validate_args(args: argparse.Namespace) -> None:
    if not args.base_url:
        raise ValueError("Informe --base-url, variável MSTR_BASE_URL ou preencha MSTR_BASE_URL no script")
    if not args.username:
        raise ValueError("Informe --username, variável MSTR_USERNAME ou preencha MSTR_USERNAME no script")
    if not args.password:
        raise ValueError("Informe --password, variável MSTR_PASSWORD ou preencha MSTR_PASSWORD no script")
    if not re.match(r"^https?://", args.base_url.strip(), flags=re.IGNORECASE):
        raise ValueError("--base-url deve iniciar com http:// ou https://")


def main() -> None:
    args = parse_args()
    validate_args(args)

    original_base_url = args.base_url.strip()
    normalized_base_url = normalize_base_url(original_base_url)

    ctx = ApiContext(
        base_url=normalized_base_url,
        timeout=args.timeout,
        verify_ssl=not args.insecure,
    )
    client = MstrClient(ctx)

    include_report_ids = parse_include_report_ids(args.include_report_ids)

    try:
        if normalized_base_url != original_base_url.rstrip("/"):
            print(f"URL base ajustada automaticamente para: {normalized_base_url}")

        print("[1/5] Autenticando...")
        client.login(args.username, args.password)

        print("[2/5] Selecionando projeto...")
        project_id = find_project_id(client, args.project_id, args.project_name)
        client.set_project(project_id)

        print("[3/5] Buscando atributos (REST 2024 /api/searches/results?type=12)...")
        attributes = get_attributes_2024(
            client,
            workers=args.workers,
            resolve_datatype=args.resolve_attribute_datatype,
        )
        attributes_keyword = [item for item in attributes if contains_keyword(item.get("name"))]

        print("[4/5] Buscando relatórios Freeform com CPF/CNPJ (REST 2024)...")
        freeform_reports = get_freeform_reports_with_cpf_cnpj_2024(
            client,
            workers=args.workers,
            include_report_ids=include_report_ids,
        )

        enrich_keyword_attributes_from_report_columns(attributes_keyword, freeform_reports)

        print("[5/5] Salvando JSON e Excel...")
        output_payload = {
            "generated_at_utc": datetime.now(timezone.utc).isoformat(),
            "base_url": normalized_base_url,
            "project_id": project_id,
            "attributes": attributes,
            "attributes_with_cpf_or_cnpj": attributes_keyword,
            "freeform_sql_reports_with_cpf_or_cnpj": freeform_reports,
            "summary": {
                "attributes_count": len(attributes),
                "attributes_with_cpf_or_cnpj_count": len(attributes_keyword),
                "reports_count": len(freeform_reports),
            },
        }

        with open(args.output, "w", encoding="utf-8") as file:
            json.dump(output_payload, file, ensure_ascii=False, indent=2)

        save_excel(args.excel_output, attributes, attributes_keyword, freeform_reports)

        print(f"Concluído. JSON: {args.output}")
        print(f"Concluído. Excel: {args.excel_output}")
        print(
            "Resumo: "
            f"{len(attributes)} atributos totais, "
            f"{len(attributes_keyword)} atributos com CPF/CNPJ, "
            f"{len(freeform_reports)} relatórios Freeform com CPF/CNPJ."
        )
    finally:
        client.logout()


if __name__ == "__main__":
    main()
